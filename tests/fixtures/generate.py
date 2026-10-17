"""Regenerate small, synthetic document fixtures. Not required to run the app/tests.

Optional developer dependencies: reportlab, python-docx, pypdf.
All content is invented test data; no account keys or personal documents.
"""
from pathlib import Path

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

root = Path(__file__).resolve().parent
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
pdf = canvas.Canvas(str(root / "sample.pdf"), invariant=1)
pdf.setFont("Helvetica", 14)
pdf.drawString(50, 780, "FPGA document translation")
pdf.drawString(50, 752, "The accelerator improves throughput.")
pdf.setFont("STSong-Light", 14)
pdf.drawString(50, 716, "中文提取测试：定点运算与激活函数。")
pdf.showPage()
pdf.setFont("Helvetica", 12)
pdf.drawString(50, 780, "Second page: preserve complete document text.")
pdf.save()

image_pdf = canvas.Canvas(str(root / "image-only.pdf"), invariant=1)
image_pdf.rect(70, 400, 350, 250, fill=1)
image_pdf.save()

reader = PdfReader(root / "sample.pdf")
writer = PdfWriter()
writer.add_page(reader.pages[0])
writer.add_page(PdfReader(root / "image-only.pdf").pages[0])
writer.add_page(reader.pages[1])
with (root / "mixed-pages.pdf").open("wb") as output:
    writer.write(output)

protected = PdfWriter()
protected.add_page(reader.pages[0])
protected.encrypt("synthetic-test-password")
with (root / "password.pdf").open("wb") as output:
    protected.write(output)

doc = Document()
doc.core_properties.author = "MyTranslator test fixture"
doc.core_properties.last_modified_by = "MyTranslator tests"
doc.add_heading("FPGA document translation", 0)
doc.add_paragraph("The accelerator improves throughput.")
doc.add_paragraph("中文提取测试：定点运算与激活函数。")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Term"
table.cell(0, 1).text = "Translation"
table.cell(1, 0).text = "throughput"
table.cell(1, 1).text = "吞吐量"
doc.sections[0].header.paragraphs[0].text = "Synthetic header"
doc.sections[0].footer.paragraphs[0].text = "Synthetic footer"
doc.save(root / "sample.docx")

print("Created synthetic PDF and Word fixtures.")
